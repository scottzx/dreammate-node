#!/usr/bin/env python3
"""Offline artifact extraction. All stdout is one JSON result; no model imports."""
import base64
import contextlib
import datetime
import hashlib
import importlib.util
import json
import math
import mimetypes
import os
import posixpath
from pathlib import Path
import shutil
import socket
import struct
import subprocess
import sys
import tempfile
from urllib.parse import urlsplit
import xml.etree.ElementTree as ET
import zipfile

PARSER_VERSION = "artifact-local-v1"
MIMES = {"pdf": "application/pdf", "presentation": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
         "spreadsheet": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "video": "video/mp4"}


def deny_python_network(*args, **kwargs):
    raise RuntimeError("Network access is disabled in the local artifact extractor")


def network_isolation():
    sandbox = shutil.which("sandbox-exec") if sys.platform == "darwin" else None
    if sandbox:
        return {"level": "os-sandbox", "command": [sandbox, "-p", "(version 1)(allow default)(deny network*)"]}
    if sys.platform.startswith("linux"):
        unshare = shutil.which("unshare")
        if unshare:
            try:
                checked = subprocess.run([unshare, "--user", "--map-root-user", "--net", "true"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=3)
                if checked.returncode == 0:
                    return {"level": "os-network-namespace", "command": [unshare, "--user", "--map-root-user", "--net"]}
            except (OSError, subprocess.TimeoutExpired):
                pass
    return {"level": "application-only", "command": []}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def file_sha(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def chunks(text, size=200, overlap=30):
    if not isinstance(size, int) or size < 32 or not isinstance(overlap, int) or overlap < 0 or overlap >= size:
        raise ValueError("chunkChars must be >=32 and chunkOverlap must be >=0 and smaller than chunkChars")
    if not text.strip():
        return []
    result = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        result.append({"modality": "text", "text": text[start:end], "content_hash": sha(text[start:end].encode("utf-8")),
                       "text_start": start, "text_end": end})
        if end == len(text):
            break
        start = end - overlap
    return result


def video_segments(duration_ms, segment_ms=6000, overlap_ms=1000):
    if not isinstance(duration_ms, int) or duration_ms < 1:
        raise ValueError("Video duration must be a positive number of milliseconds")
    if not isinstance(segment_ms, int) or not 1000 <= segment_ms <= 6000:
        raise ValueError("segmentMs must be between 1000 and 6000")
    if not isinstance(overlap_ms, int) or not 0 <= overlap_ms < segment_ms:
        raise ValueError("overlapMs must be >=0 and smaller than segmentMs")
    result = []
    start = 0
    while start < duration_ms:
        end = min(start + segment_ms, duration_ms)
        result.append((start, end))
        if end == duration_ms:
            break
        start = end - overlap_ms
    assert result[0][0] == 0 and result[-1][1] == duration_ms
    assert all(result[index][0] <= result[index - 1][1] for index in range(1, len(result)))
    return result


class Extractor:
    def __init__(self, source, cache, options):
        self.source = Path(source).resolve(strict=True)
        self.cache = Path(cache).resolve()
        self.cache.mkdir(parents=True, exist_ok=True)
        self.options = options
        self.warnings = []
        self.cache_bytes = 0
        self.isolation = network_isolation()
        if self.isolation["level"] == "application-only":
            self.warnings.append("OS network isolation unavailable: Python outbound connections are blocked, Office external relationships are removed, and FFmpeg accepts file/pipe only; external converter isolation is not enforced by the OS")

    def limit(self, name, default, minimum=1):
        value = self.options.get(name, default)
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < minimum:
            raise ValueError(f"{name} must be >= {minimum}")
        return value

    def budget(self, value, name, default):
        if value > self.limit(name, default):
            raise ValueError(f"Artifact exceeds {name} budget ({value})")

    def output_path(self, name):
        """Check before handing a writable destination to an external tool."""
        path = self.cache / name
        if not path.resolve().is_relative_to(self.cache):
            raise ValueError("Artifact output escaped its cache directory")
        if path.is_symlink():
            raise ValueError("Refusing a symbolic link as an artifact output")
        if path.exists() and path.samefile(self.source):
            raise ValueError("Refusing to overwrite the original artifact")
        return path

    def command(self, args, timeout=120):
        result = subprocess.run(self.isolation["command"] + [str(arg) for arg in args], stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                timeout=min(timeout, self.limit("commandTimeoutSeconds", 180)), check=False)
        if result.returncode:
            raise RuntimeError(f"{Path(args[0]).name} failed: {result.stderr.decode('utf-8', errors='replace')[-4000:]}")
        return result.stdout

    def text_representations(self, text, locator):
        self.budget(len(text), "maxUnitChars", 2_000_000)
        result = chunks(text, int(self.limit("chunkChars", 200, 32)), int(self.limit("chunkOverlap", 30, 0)))
        for item in result:
            item["locator"] = locator
        return result

    def media(self, path, modality):
        path = Path(path).resolve(strict=True)
        if not path.is_relative_to(self.cache):
            raise ValueError("Rendered artifact escaped its cache directory")
        self.cache_bytes += path.stat().st_size
        self.budget(self.cache_bytes, "maxCacheBytes", 1024 * 1024 * 1024)
        return {"modality": modality, "path": str(path), "content_hash": file_sha(path)}

    def render_pdf(self, path, page_count, prefix="page"):
        command = shutil.which("pdftoppm")
        if not command:
            self.warnings.append("Visual extraction unavailable: pdftoppm is not installed; PDF pages retain text and locators only")
            return {}
        dpi = self.limit("renderDpi", 100, 36)
        if dpi > 300:
            raise ValueError("renderDpi must be <=300")
        result = {}
        for page in range(1, page_count + 1):
            image = self.output_path(f"{prefix}-{page:05d}.png")
            stem = image.with_suffix("")
            self.command([command, "-f", page, "-l", page, "-r", dpi, "-singlefile", "-scale-to", 2400,
                          "-png", str(path), str(stem)])
            result[page] = self.media(image, "image")
        return result

    def pdf_pages(self, path):
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("Encrypted PDF requires a local decrypted copy")
        self.budget(len(reader.pages), "maxPages", 1000)
        texts = None
        command = shutil.which("pdftotext")
        if command:
            output = self.command([command, "-layout", "-enc", "UTF-8", str(path), "-"])
            pages = output.decode("utf-8").split("\f")
            if pages and not pages[-1].strip():
                pages.pop()
            if len(pages) == len(reader.pages):
                texts = pages
            else:
                self.warnings.append("pdftotext page count mismatch; all pages extracted with pypdf instead")
        if texts is None:
            texts = [page.extract_text() or "" for page in reader.pages]
        return texts

    def pdf(self):
        texts = self.pdf_pages(self.source)
        images = self.render_pdf(self.source, len(texts)) if self.options.get("visual", True) else {}
        units = []
        for page, text in enumerate(texts, 1):
            locator = {"page": page}
            representations = self.text_representations(text, locator)
            if page in images:
                representations.append(images[page])
            if not text.strip():
                self.warnings.append(f"Page {page} has no extractable text; visual representation is required for retrieval (OCR was not run)")
            units.append({"kind": "page", "locator": locator, "text": text, "representations": representations})
        return self.finish("pdf", units)

    def clean_presentation_links(self):
        if self.source.suffix.lower() not in (".pptx", ".pptm"):
            return self.source
        cleaned = self.output_path("local-source" + self.source.suffix.lower())
        removed = 0
        ET.register_namespace("", "http://schemas.openxmlformats.org/package/2006/relationships")
        with zipfile.ZipFile(self.source) as original, zipfile.ZipFile(cleaned, "w", compression=zipfile.ZIP_DEFLATED) as result:
            self.budget(sum(info.file_size for info in original.infolist()), "maxArchiveBytes", 512 * 1024 * 1024)
            for info in original.infolist():
                data = original.read(info.filename)
                if info.filename.endswith(".rels"):
                    root = ET.fromstring(data)
                    changed = False
                    for relation in list(root):
                        target = relation.get("Target", "")
                        # Reject URI/absolute/traversing targets as well as the
                        # normal OOXML External flag; malformed rels cannot opt
                        # back into a remote resource by omitting TargetMode.
                        package_base = posixpath.dirname(posixpath.dirname(info.filename))
                        normalized = posixpath.normpath(posixpath.join(package_base, target.replace("\\", "/")))
                        external = (relation.get("TargetMode", "").strip().lower() == "external"
                                    or bool(urlsplit(target).scheme) or target.startswith(("/", "\\"))
                                    or normalized == ".." or normalized.startswith("../"))
                        if external:
                            root.remove(relation)
                            removed += 1
                            changed = True
                    if changed:
                        data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                result.writestr(info, data)
        if removed:
            self.warnings.append(f"Removed {removed} external Office relationships from a cached rendering copy; original file was preserved and external resources were not fetched")
        return cleaned

    def convert_presentation(self):
        command = shutil.which("soffice") or shutil.which("libreoffice")
        if not command:
            return None
        local_source = self.clean_presentation_links()
        out = self.output_path("office")
        out.mkdir(exist_ok=True)
        target = self.output_path(Path("office") / (local_source.stem + ".pdf"))
        # A failed/no-output conversion must never reuse an older cached PDF.
        target.unlink(missing_ok=True)
        # Disposable profile prevents touching user settings and disables macros.
        with tempfile.TemporaryDirectory(prefix="artifact-office-") as directory:
            profile = Path(directory)
            (profile / "user").mkdir()
            (profile / "user" / "registrymodifications.xcu").write_text(
                '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
                '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item>'
                '</oor:items>', encoding="utf-8")
            self.command([command, "-env:UserInstallation=" + profile.as_uri(), "--headless", "--nologo", "--nodefault",
                          "--norestore", "--convert-to", "pdf:impress_pdf_Export", "--outdir", str(out), str(local_source)], timeout=180)
        if not target.is_file():
            raise RuntimeError("LibreOffice did not produce a PDF")
        return target

    def presentation(self):
        slides = None
        if self.source.suffix.lower() in (".pptx", ".pptm"):
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            presentation = Presentation(str(self.source))
            self.budget(len(presentation.slides), "maxPages", 1000)
            slides = []
            def shape_text(shape):
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    return "\n".join(shape_text(child) for child in shape.shapes)
                if shape.has_table:
                    return "\n".join("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
                return shape.text if shape.has_text_frame else ""
            for slide in presentation.slides:
                body = "\n".join(filter(None, (shape_text(shape) for shape in slide.shapes)))
                notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ""
                slides.append((body, notes))
        pdf = None
        conversion_failed = False
        if self.options.get("visual", True) or slides is None:
            try:
                pdf = self.convert_presentation()
            except Exception as error:
                if slides is None:
                    raise
                self.warnings.append(f"Presentation visual extraction failed: {error}")
                conversion_failed = True
            if pdf is None and not conversion_failed:
                if slides is None:
                    raise RuntimeError("Legacy PPT requires local LibreOffice (soffice); no file content was indexed")
                self.warnings.append("Presentation visual extraction unavailable: local LibreOffice (soffice) is not installed")
        if pdf:
            pdf_texts = self.pdf_pages(pdf)
            if slides is None:
                slides = [(text, "") for text in pdf_texts]
                self.warnings.append("Legacy PPT text was extracted from rendered slides; speaker notes are unavailable")
            if len(pdf_texts) != len(slides):
                raise ValueError(f"Rendered slide count mismatch ({len(pdf_texts)} vs {len(slides)}); refusing incomplete indexing")
        images = self.render_pdf(pdf, len(slides), "slide") if pdf and self.options.get("visual", True) else {}
        units = []
        for index, (body, notes) in enumerate(slides, 1):
            text = body + (("\nSpeaker notes:\n" + notes) if notes else "")
            locator = {"slide": index}
            reps = self.text_representations(text, locator)
            if index in images:
                reps.append(images[index])
            units.append({"kind": "slide", "locator": locator, "text": text, "notes": notes, "representations": reps})
        return self.finish("presentation", units)

    @staticmethod
    def json_value(value):
        if isinstance(value, (datetime.date, datetime.time, datetime.datetime)):
            return value.isoformat()
        if isinstance(value, float) and not math.isfinite(value):
            return str(value)
        return value

    def spreadsheet(self):
        from openpyxl.utils.cell import get_column_letter
        row_block = int(self.limit("rowsPerUnit", 20))
        sheets = []
        total_cells = 0
        if self.source.suffix.lower() == ".xls":
            import xlrd
            from xlrd.formula import decompile_formula, FMLA_TYPE_CELL
            workbook = xlrd.open_workbook(str(self.source), on_demand=True, logfile=sys.stderr)
            try:
                for sheet_index in range(workbook.nsheets):
                    sheet = workbook.sheet_by_index(sheet_index)
                    total_cells += sheet.nrows * sheet.ncols
                    self.budget(total_cells, "maxCells", 1_000_000)
                    formulas = {}
                    # Preserve original BIFF formula bytecode, even when a legacy
                    # expression cannot be decompiled. Never evaluate formulas.
                    position = workbook._sh_abs_posn[sheet_index]
                    memory = workbook.mem
                    while position + 4 <= len(memory):
                        code, length = struct.unpack_from("<HH", memory, position)
                        data = memory[position + 4:position + 4 + length]
                        position += length + 4
                        if code == 0x000A:
                            break
                        if code in (0x0006, 0x0206, 0x0406) and len(data) >= 22:
                            row, col = struct.unpack_from("<HH", data)
                            count = struct.unpack_from("<H", data, 20)[0]
                            tokens = bytes(data[22:22 + count])
                            formula = {"formula_biff_base64": base64.b64encode(tokens).decode("ascii")}
                            try:
                                expression = decompile_formula(workbook, tokens, len(tokens), FMLA_TYPE_CELL, browx=row, bcolx=col)
                                if expression:
                                    formula["formula"] = "=" + expression
                            except Exception:
                                pass
                            if "formula" not in formula:
                                self.warnings.append(f"XLS formula at {sheet.name}!{get_column_letter(col+1)}{row+1} retained as original BIFF tokens; expression unavailable")
                            formulas[(row, col)] = formula
                    rows = []
                    for row in range(sheet.nrows):
                        cells = []
                        for col in range(sheet.ncols):
                            cell = sheet.cell(row, col)
                            if cell.ctype == xlrd.XL_CELL_EMPTY and (row, col) not in formulas:
                                continue
                            value = cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                try:
                                    value = xlrd.xldate_as_datetime(value, workbook.datemode).isoformat()
                                except ValueError:
                                    pass
                            item = {"coordinate": f"{get_column_letter(col + 1)}{row + 1}", "value": self.json_value(value), "data_type": str(cell.ctype)}
                            if (row, col) in formulas:
                                item.update(formulas[(row, col)])
                                item["cached_value"] = item["value"]
                            cells.append(item)
                        rows.append(cells)
                    sheets.append((sheet.name, rows, sheet.ncols, bool(sheet.visibility), []))
            finally:
                workbook.release_resources()
        else:
            import openpyxl
            with zipfile.ZipFile(self.source) as archive:
                self.budget(sum(info.file_size for info in archive.infolist()), "maxArchiveBytes", 512 * 1024 * 1024)
            workbook = openpyxl.load_workbook(self.source, read_only=True, data_only=False, keep_links=False)
            cached = openpyxl.load_workbook(self.source, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook.worksheets:
                    # Some producers save incorrect <dimension> bounds. Stream
                    # every actual row rather than silently losing trailing cells.
                    sheet.reset_dimensions()
                    cached_sheet = cached[sheet.title]
                    cached_sheet.reset_dimensions()
                    rows = []
                    width = 0
                    for row, cached_row in zip(sheet.iter_rows(), cached_sheet.iter_rows()):
                        total_cells += len(row)
                        width = max(width, len(row))
                        self.budget(total_cells, "maxCells", 1_000_000)
                        cells = []
                        for cell, cache_cell in zip(row, cached_row):
                            if cell.value is None:
                                continue
                            value = self.json_value(cell.value.text if cell.data_type == "f" and hasattr(cell.value, "text") else cell.value)
                            item = {"coordinate": cell.coordinate, "value": value, "data_type": cell.data_type, "number_format": cell.number_format}
                            if cell.data_type == "f":
                                item["formula"] = value
                                if hasattr(cell.value, "ref"):
                                    item["formula_range"] = cell.value.ref
                                item["cached_value"] = self.json_value(cache_cell.value)
                                if cache_cell.value is None:
                                    self.warnings.append(f"Formula {sheet.title}!{cell.coordinate} has no saved cached value; no calculation was performed")
                            cells.append(item)
                        rows.append(cells)
                    sheets.append((sheet.title, rows, width, sheet.sheet_state != "visible", []))
            finally:
                workbook.close()
                cached.close()
        units = []
        for name, rows, width, hidden, merged in sheets:
            header = next((row for row in rows if row), [])
            for start in range(0, max(1, len(rows)), row_block):
                end = min(start + row_block, max(1, len(rows)))
                cell_range = f"A{start + 1}:{get_column_letter(max(width, 1))}{end}"
                locator = {"sheet": name, "range": cell_range}
                cells = [cell for row in rows[start:end] for cell in row]
                def describe(cell):
                    result = f"{cell['coordinate']}: {cell['value']}"
                    if "formula" in cell:
                        result += f" (formula={cell['formula']}; cached_value={cell.get('cached_value')})"
                    elif "formula_biff_base64" in cell:
                        result += f" (formula_biff_base64={cell['formula_biff_base64']}; cached_value={cell.get('cached_value')})"
                    return result
                text = f"Worksheet: {name}\nRange: {cell_range}\nHeaders: " + " | ".join(describe(cell) for cell in header) + "\n"
                text += "\n".join(describe(cell) for cell in cells)
                units.append({"kind": "cell_range", "locator": locator, "text": text, "cells": cells,
                              "header_cells": header, "hidden": hidden, "representations": self.text_representations(text, locator)})
        return self.finish("spreadsheet", units)

    def video(self):
        probe = shutil.which("ffprobe")
        if not probe:
            raise RuntimeError("Video extraction requires local ffprobe")
        details = json.loads(self.command([probe, "-protocol_whitelist", "file,pipe", "-v", "error", "-show_format", "-show_streams", "-of", "json", str(self.source)]))
        stream = next((item for item in details.get("streams", []) if item.get("codec_type") == "video" and not item.get("disposition", {}).get("attached_pic")), None)
        if stream is None:
            raise ValueError("Artifact has no playable video stream")
        duration = float(stream.get("duration") or details.get("format", {}).get("duration") or 0)
        if not math.isfinite(duration) or duration <= 0:
            raise ValueError("Video duration is unavailable or invalid")
        self.budget(duration, "maxVideoSeconds", 36_000)
        duration_ms = math.ceil(duration * 1000)
        segments = video_segments(duration_ms, int(self.limit("segmentMs", 6000)), int(self.limit("overlapMs", 1000, 0)))
        self.budget(len(segments), "maxUnits", 20_000)
        command = shutil.which("ffmpeg")
        if self.options.get("visual", True) and not command:
            raise RuntimeError("Video extraction requires local ffmpeg")
        fps = self.limit("videoFps", 2)
        dimension = int(self.limit("videoMaxDimension", 720, 64))
        if fps > 30 or dimension > 2160:
            raise ValueError("videoFps must be <=30 and videoMaxDimension must be <=2160")
        units = []
        for index, (start, end) in enumerate(segments):
            representations = []
            if self.options.get("visual", True):
                path = self.output_path(f"segment-{index+1:05d}-{start}-{end}.mp4")
                self.command([command, "-nostdin", "-protocol_whitelist", "file,pipe", "-v", "error", "-y", "-ss", f"{start/1000:.3f}", "-i", str(self.source),
                              "-t", f"{(end-start)/1000:.3f}", "-map", f"0:{stream['index']}", "-an", "-sn", "-dn",
                              "-vf", f"fps={fps}:round=up,scale={dimension}:{dimension}:force_original_aspect_ratio=decrease:force_divisible_by=2",
                              "-c:v", "libx264", "-preset", "veryfast", "-crf", "22", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(path)])
                if not path.is_file() or path.stat().st_size < 100:
                    raise RuntimeError(f"Video segment {start}-{end} did not render; refusing partial time coverage")
                rendered = json.loads(self.command([probe, "-protocol_whitelist", "file,pipe", "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=nb_frames,duration", "-of", "json", str(path)]))
                if not rendered.get("streams") or int(rendered["streams"][0].get("nb_frames", "0")) < 1:
                    raise RuntimeError(f"Video segment {start}-{end} contains no frames; refusing partial time coverage")
                representations.append(self.media(path, "video"))
            locator = {"start_ms": start, "end_ms": end}
            units.append({"kind": "video_segment", "locator": locator, "text": "", "representations": representations})
        self.warnings.append("Video indexing covers visual frames only; audio and speech were not transcribed")
        if not self.options.get("visual", True):
            self.warnings.append("Visual extraction disabled: video has time locators but no searchable media representations")
        result = self.finish("video", units)
        result["duration_ms"] = duration_ms
        result["sampling"] = {"segment_ms": self.options.get("segmentMs", 6000), "overlap_ms": self.options.get("overlapMs", 1000), "fps": fps, "max_dimension": dimension}
        return result

    def finish(self, kind, units):
        self.budget(len(units), "maxUnits", 20_000)
        self.budget(sum(len(unit.get("text", "")) for unit in units), "maxTextChars", 32_000_000)
        return {"kind": kind, "mime": mimetypes.guess_type(str(self.source))[0] or MIMES[kind], "title": self.source.stem, "parser_version": PARSER_VERSION,
                "units": units, "warnings": list(dict.fromkeys(self.warnings)), "offline_isolation": self.isolation["level"]}

    def extract(self):
        extension = self.source.suffix.lower()
        if extension in (".pptx", ".pptm", ".xlsx", ".xlsm"):
            with zipfile.ZipFile(self.source) as archive:
                self.budget(sum(info.file_size for info in archive.infolist()), "maxArchiveBytes", 512 * 1024 * 1024)
        if extension == ".pdf":
            return self.pdf()
        if extension in (".pptx", ".ppt", ".pptm"):
            return self.presentation()
        if extension in (".xlsx", ".xls", ".xlsm"):
            return self.spreadsheet()
        if extension in (".mp4", ".mov", ".webm", ".mkv", ".avi", ".m4v", ".mpeg", ".mpg"):
            return self.video()
        raise ValueError(f"Unsupported artifact extension: {extension or '(none)'}")


def doctor():
    modules = {name: importlib.util.find_spec(name) is not None for name in ("pypdf", "pptx", "openpyxl", "xlrd", "PIL")}
    commands = {name: shutil.which(name) for name in ("pdftoppm", "pdftotext", "soffice", "libreoffice", "ffmpeg", "ffprobe")}
    return {"python": sys.executable, "parser_version": PARSER_VERSION, "modules": modules, "commands": commands,
            "offline_isolation": network_isolation()["level"],
            "capabilities": {"pdf_text": modules["pypdf"], "pdf_visual": modules["pypdf"] and bool(commands["pdftoppm"]),
                             "pptx_text": modules["pptx"], "presentation_visual": bool(commands["soffice"] or commands["libreoffice"]) and bool(commands["pdftoppm"]),
                             "xlsx": modules["openpyxl"], "xls": modules["xlrd"] and modules["openpyxl"], "video": bool(commands["ffmpeg"] and commands["ffprobe"])}}


def main():
    socket.socket.connect = deny_python_network
    socket.socket.connect_ex = deny_python_network
    socket.socket.sendto = deny_python_network
    if hasattr(socket.socket, "sendmsg"):
        socket.socket.sendmsg = deny_python_network
    socket.create_connection = deny_python_network
    socket.getaddrinfo = deny_python_network
    request = json.load(sys.stdin)
    # Third-party parsers sometimes print diagnostics. Keep protocol stdout clean.
    with contextlib.redirect_stdout(sys.stderr):
        result = doctor() if request.get("operation") == "doctor" else Extractor(request["file"], request["cacheDir"], request.get("options", {})).extract()
    print(json.dumps(result, ensure_ascii=False, allow_nan=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"{type(error).__name__}: {error}", file=sys.stderr)
        sys.exit(1)
