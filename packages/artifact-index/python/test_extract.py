import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
import zipfile

from extract import Extractor, chunks, video_segments, file_sha


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def extractor(self, source, **options):
        return Extractor(source, self.root / "cache", options)

    def test_chunks_cover_full_unicode_text_with_bounded_overlap(self):
        text = "眼镜木马🔎" * 117
        output = chunks(text)
        self.assertTrue(all(len(item["text"]) <= 200 for item in output))
        rebuilt = output[0]["text"] + "".join(item["text"][30:] for item in output[1:])
        self.assertEqual(rebuilt, text)
        self.assertEqual(output[-1]["text_end"], len(text))

    def test_segments_cover_exact_tail_and_no_gaps(self):
        for duration in [1, 542, 6000, 6001, 12542, 18542, 60000]:
            with self.subTest(duration=duration):
                output = video_segments(duration)
                self.assertEqual(output[0][0], 0)
                self.assertEqual(output[-1][1], duration)
                self.assertTrue(all(end - start <= 6000 for start, end in output))
                self.assertTrue(all(output[i][0] <= output[i-1][1] for i in range(1, len(output))))
        self.assertEqual(video_segments(18542), [(0, 6000), (5000, 11000), (10000, 16000), (15000, 18542)])
        with self.assertRaises(ValueError):
            video_segments(10, 6000, 6000)

    def test_outputs_cannot_overwrite_original_or_escape_through_links(self):
        source = self.root / "source.pdf"
        source.write_bytes(b"original bytes must stay intact")
        extractor = self.extractor(source)
        before = file_sha(source)
        hard_link = extractor.cache / "page-00001.png"
        os.link(source, hard_link)
        with self.assertRaisesRegex(ValueError, "original artifact"):
            extractor.output_path(hard_link.name)
        hard_link.unlink()
        symbolic = extractor.cache / "page-00001.png"
        symbolic.symlink_to(source)
        with self.assertRaisesRegex(ValueError, "escaped|symbolic"):
            extractor.output_path(symbolic.name)
        symbolic.unlink()
        outside = self.root / "outside"
        outside.mkdir()
        (extractor.cache / "office").symlink_to(outside, target_is_directory=True)
        with self.assertRaisesRegex(ValueError, "escaped|symbolic"):
            extractor.output_path("office/output.pdf")
        self.assertEqual(before, file_sha(source))

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf unavailable")
    def test_pdf_keeps_all_pages_including_scanned_blank_page(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        path = self.root / "sample.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(400, 300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 18 Tf 20 200 Td (Local PDF glasses document) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.add_blank_page(400, 300)
        writer.write(path)
        output = self.extractor(path, visual=bool(shutil.which("pdftoppm"))).extract()
        self.assertEqual(len(output["units"]), 2)
        self.assertEqual([item["locator"] for item in output["units"]], [{"page": 1}, {"page": 2}])
        self.assertIn("glasses", output["units"][0]["text"])
        if shutil.which("pdftoppm"):
            self.assertTrue(all(any(rep["modality"] == "image" for rep in item["representations"]) for item in output["units"]))
        self.assertTrue(any("Page 2" in warning for warning in output["warnings"]))
        with self.assertRaisesRegex(ValueError, "maxPages"):
            self.extractor(path, maxPages=1, visual=False).extract()

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl unavailable")
    def test_spreadsheet_ranges_headers_formulas_and_saved_values(self):
        from openpyxl import Workbook
        path = self.root / "budget.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "预算"
        sheet.append(["产品", "数量", "单价", "总额"])
        for index in range(1, 26):
            sheet.append([f"眼镜{index}", index, 10, f"=B{index+1}*C{index+1}"])
        book.save(path)
        # Supply a cached value exactly as a real spreadsheet application saves it.
        with zipfile.ZipFile(path) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        files["xl/worksheets/sheet1.xml"] = files["xl/worksheets/sheet1.xml"].replace(b"<f>B2*C2</f><v></v>", b"<f>B2*C2</f><v>10</v>")
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in files.items():
                archive.writestr(name, data)
        output = self.extractor(path, rowsPerUnit=20).extract()
        self.assertEqual([unit["locator"] for unit in output["units"]], [{"sheet": "预算", "range": "A1:D20"}, {"sheet": "预算", "range": "A21:D26"}])
        formula = next(cell for cell in output["units"][0]["cells"] if cell["coordinate"] == "D2")
        self.assertEqual(formula["formula"], "=B2*C2")
        self.assertEqual(formula["cached_value"], 10)
        self.assertIn("产品", output["units"][1]["text"])
        self.assertIn("眼镜25", output["units"][-1]["text"])
        self.assertTrue(all(len(rep["text"]) <= 200 for unit in output["units"] for rep in unit["representations"]))
        with self.assertRaisesRegex(ValueError, "maxCells"):
            self.extractor(path, maxCells=10).extract()

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx unavailable")
    def test_pptx_preserves_slides_and_speaker_notes(self):
        from pptx import Presentation
        path = self.root / "story.pptx"
        presentation = Presentation()
        for text in ["Glasses slide", "Trojan horse slide"]:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = text
            slide.placeholders[1].text = "Local extraction test"
            slide.notes_slide.notes_text_frame.text = "Private speaker notes " + text
        presentation.save(path)
        output = self.extractor(path, visual=False).extract()
        self.assertEqual([unit["locator"] for unit in output["units"]], [{"slide": 1}, {"slide": 2}])
        self.assertIn("Private speaker notes", output["units"][1]["text"])
        if (shutil.which("soffice") or shutil.which("libreoffice")) and shutil.which("pdftoppm"):
            visual = self.extractor(path, visual=True).extract()
            self.assertTrue(all(any(rep["modality"] == "image" for rep in unit["representations"]) for unit in visual["units"]))

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx unavailable")
    def test_external_office_links_removed_without_changing_original(self):
        from pptx import Presentation
        import xml.etree.ElementTree as ET
        path = self.root / "external.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        run = slide.shapes.title.text_frame.paragraphs[0].add_run()
        run.text = "External destination"
        run.hyperlink.address = "https://example.invalid/must-not-fetch"
        presentation.save(path)
        # A malformed producer may omit TargetMode=External. URI targets must
        # still be stripped before a converter sees the rendering copy.
        with zipfile.ZipFile(path) as archive:
            original_parts = {name: archive.read(name) for name in archive.namelist()}
        rels_name = "ppt/slides/_rels/slide1.xml.rels"
        rels = ET.fromstring(original_parts[rels_name])
        ET.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
            "Id": "rIdMalformedRemote", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
            "Target": "https://example.invalid/unflagged-image.png"})
        original_parts[rels_name] = ET.tostring(rels)
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in original_parts.items():
                archive.writestr(name, data)
        before = file_sha(path)
        extractor = self.extractor(path)
        cleaned = extractor.clean_presentation_links()
        with zipfile.ZipFile(cleaned) as archive:
            for name in archive.namelist():
                if name.endswith(".rels"):
                    self.assertTrue(all(item.get("TargetMode", "").lower() != "external" and not item.get("Target", "").startswith("https:") for item in ET.fromstring(archive.read(name))))
        self.assertEqual(before, file_sha(path))
        self.assertTrue(any("external Office relationships" in warning for warning in extractor.warnings))

    @unittest.skipUnless(importlib.util.find_spec("openpyxl") and importlib.util.find_spec("xlrd") and (shutil.which("soffice") or shutil.which("libreoffice")), "legacy XLS fixture dependencies unavailable")
    def test_legacy_xls_preserves_formula_bytecode_and_cached_value(self):
        from openpyxl import Workbook
        source = self.root / "legacy-source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Legacy budget"
        sheet.append(["Quantity", "Price", "Amount"])
        sheet.append([3, 7, "=A2*B2"])
        workbook.save(source)
        profile = self.root / "xls-fixture-profile"
        out = self.root / "xls-fixture"
        out.mkdir()
        extractor = self.extractor(source)
        extractor.command([shutil.which("soffice") or shutil.which("libreoffice"), "-env:UserInstallation=" + profile.as_uri(), "--headless", "--convert-to", "xls:MS Excel 97", "--outdir", str(out), str(source)])
        output = self.extractor(out / "legacy-source.xls").extract()
        formula = next(cell for cell in output["units"][0]["cells"] if cell["coordinate"] == "C2")
        self.assertTrue(formula["formula_biff_base64"])
        self.assertEqual(formula["cached_value"], 21)
        self.assertIn("A2", formula["formula"])
        self.assertEqual(output["mime"], "application/vnd.ms-excel")

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "ffmpeg unavailable")
    def test_video_extracts_real_playable_segments_with_tail(self):
        path = self.root / "short.mp4"
        subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-f", "lavfi", "-i", "color=c=red:s=160x120:r=25:d=7.543", "-c:v", "libx264", "-pix_fmt", "yuv420p", str(path)], check=True)
        output = self.extractor(path).extract()
        self.assertEqual(len(output["units"]), 2)
        self.assertEqual(output["units"][-1]["locator"]["end_ms"], output["duration_ms"])
        self.assertGreater(output["duration_ms"], 7500)
        self.assertLess(output["duration_ms"], 7600)
        self.assertEqual(output["units"][1]["locator"]["start_ms"], 5000)
        for unit in output["units"]:
            rep = unit["representations"][0]
            self.assertEqual(rep["modality"], "video")
            self.assertGreater(Path(rep["path"]).stat().st_size, 100)
        with self.assertRaisesRegex(ValueError, "maxVideoSeconds"):
            self.extractor(path, maxVideoSeconds=1).extract()


if __name__ == "__main__":
    unittest.main()
